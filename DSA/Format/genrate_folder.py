import os
from openpyxl import load_workbook

# -------- Paths --------
BASE_DSA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
EXCEL_FILE = os.path.join(os.path.dirname(__file__), "DQ15.xlsx")

# Load Excel
wb = load_workbook(EXCEL_FILE)
sheet = wb.active

for row in sheet.iter_rows(min_row=2, values_only=True):
    problem_no, title, description, example, constraints = row

    if not problem_no or not title:
        continue

    folder_name = f"{int(problem_no):03d}. {title}"
    folder_path = os.path.join(BASE_DSA_DIR, folder_name)

    # ---------- Folder ----------
    if not os.path.exists(folder_path):
        os.makedirs(folder_path)
        print(f"Created folder: {folder_path}")
    else:
        print(f"Folder exists: {folder_path}")

    # ---------- code.py ----------
    code_path = os.path.join(folder_path, "code.py")
    if not os.path.exists(code_path):
        with open(code_path, "w", encoding="utf-8") as f:
            f.write(
                f"# {title}\n"
                f"# Description: {description}\n"
                f"# Example: {example}\n"
                f"# Constraints: {constraints}\n\n"
                f"def solution():\n"
                f"    pass\n"
            )
        print("Created code.py")

    # ---------- PS.md ----------
    ps_path = os.path.join(folder_path, "PS.md")
    if not os.path.exists(ps_path):
        with open(ps_path, "w", encoding="utf-8") as f:
            f.write(
                f"# {problem_no}. {title}\n\n"
                f"## Description\n{description}\n\n"
                f"## Example\n{example}\n\n"
                f"## Constraints\n{constraints}\n"
            )
        print("Created PS.md")

print("\nDSA folders created inside DSA/ safely.")
