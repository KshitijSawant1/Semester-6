import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

def generate_excel_from_csv(csv_file, output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Question Paper"

    # Header
    headers = ["Sr", "Question", "Marks"]
    ws.append(headers)

    # Style header
    for col in range(1, 4):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 2

    with open(csv_file, mode='r', encoding='utf-8') as file:
        reader = csv.DictReader(file)

        for i, row in enumerate(reader, start=1):
            question = row["Question"]
            marks = row["Marks"]

            # Question row
            ws.cell(row=row_num, column=1, value=i)
            ws.cell(row=row_num, column=2, value=question)
            ws.cell(row=row_num, column=3, value=marks)

            ws.cell(row=row_num, column=2).alignment = Alignment(wrap_text=True)

            row_num += 1

            # Blank answer row
            ws.cell(row=row_num, column=2, value="")
            ws.cell(row=row_num, column=2).alignment = Alignment(wrap_text=True)

            row_num += 1

    # Adjust column width
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 80
    ws.column_dimensions['C'].width = 10

    wb.save(output_file)
    print(f"Excel file generated: {output_file}")

# Dynamic path setup
current_dir = os.path.dirname(os.path.abspath(__file__))
csv_path = os.path.join(current_dir, "questions.csv")
excel_path = os.path.join(current_dir, "questions.xlsx")

generate_excel_from_csv(csv_path, excel_path)